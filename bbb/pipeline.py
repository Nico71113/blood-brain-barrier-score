"""Preserve input rows, score supplied descriptors, and optionally depict SMILES."""
import base64
import binascii
import csv
import io
import zipfile
from pathlib import Path
from .chemistry import molecule_data
from .scoring import FIELDS, number, score_descriptors
from .structures import lookup_structure

ALIASES = {
    "compound_id": ("compound_id", "compound ID", "compund ID", "name", "id"),
    "smiles": ("SMILES", "canonical_smiles"),
    "bbb_score_input": ("BBB SCORE", "bbb_score_input"),
    "AroR": ("AroR", "Aro_R", "AR", "Number of Aromatic Rings (Aro_R)"),
    "HA": ("HA", "Heavy_AT", "Number of Heavy Atoms (HA)"),
    "MW": ("MW", "Molecular Weight (MW)"),
    "HBA": ("HBA", "Number of Hydrogen Bond Acceptor (HBA)"),
    "HBD": ("HBD", "Number of Hydrogen Bond Donor (HBD)"),
    "TPSA": ("TPSA", "Topological Polar Surface Area(TPSA)"),
    "pKa": ("pKa", "pka_value"),
}


def normalize(value):
    return "".join(c for c in str(value).lower() if c.isalnum())


def canonical_row(row):
    if not isinstance(row, dict):
        raise ValueError("Each row must be a JSON object.")
    raw = {normalize(k): v for k, v in row.items() if k is not None}
    output = {}
    for key, aliases in ALIASES.items():
        output[key] = next((raw[normalize(a)] for a in aliases
                            if normalize(a) in raw and raw[normalize(a)] is not None
                            and str(raw[normalize(a)]).strip() != ""), None)
    output["neutral"] = row.get("neutral") is True
    return output


def analyze(row, row_id=1, draw=True):
    input_row = canonical_row(row)
    compound_id = input_row["compound_id"]
    smiles = str(input_row["smiles"] or "").strip()
    association = lookup_structure(input_row) if not smiles else None
    if association:
        smiles = association["smiles"]
    result = {"row_id": row_id, "compound_id": str(compound_id) if compound_id is not None else f"Molecule {row_id}",
              "input_smiles": smiles, "bbb_score_input": input_row["bbb_score_input"],
              "bbb_score": None, "score_status": "missing_input", "score_error": None,
              "structure_status": "missing_structure", "structure_error": None,
              "structure": None, "descriptors": {}, "contributions": [],
              "warnings": [], "input": input_row, "pka_source": "provided",
              "structure_origin": "local_association" if association else "input_smiles" if smiles else "unresolved",
              "structure_source": association["source_title"] if association else None,
              "structure_source_url": association["source_url"] if association else None,
              "structure_source_locator": association["source_locator"] if association else None,
              "structure_match_method": association["match_method"] if association else None}
    structure = None
    if smiles:
        try:
            structure = molecule_data(smiles, draw)
            result.update(structure=structure, structure_status="ok")
        except ValueError as error:
            result.update(structure_status="invalid_structure", structure_error=str(error))
    provided = {key: input_row[key] for key in FIELDS if key != "pKa"}
    use_table = any(value is not None for value in provided.values())
    if association and use_table:
        disagreements = [key for key, check in association["checks"].items() if not check["matches"]]
        if disagreements:
            result["warnings"].append("Verified structure; supplied " + ", ".join(disagreements) + " differ from RDKit. Scoring uses supplied descriptors.")
    descriptors = provided if use_table else dict(structure["descriptors"]) if structure else {}
    result["descriptor_source"] = "uploaded_or_entered" if use_table else "rdkit"
    pka = input_row["pKa"]
    if input_row["neutral"]:
        if pka is not None:
            result.update(score_status="invalid_input", score_error="Choose a pKa value or the neutral convention.")
            return result
        pka = 8.81
        result["pka_source"] = "user_confirmed_neutral_convention"
        result["warnings"].append("Neutral convention: pKa input = 8.81 (not predicted).")
    descriptors["pKa"] = pka
    result["descriptors"] = descriptors
    if structure and use_table:
        result["warnings"].append("Scoring uses supplied descriptors; the structure uses SMILES. Toolkit definitions may differ.")
    try:
        result.update(score_descriptors(descriptors))
        result["score_status"] = "ok"
    except ValueError as error:
        missing = any(descriptors.get(key) is None or str(descriptors.get(key)).strip() == "" for key in FIELDS)
        result.update(score_status="missing_input" if missing else "invalid_input", score_error=str(error))
    if not smiles:
        result["structure_error"] = "SMILES required for 2D depiction."
    return result


def read_file(filename, data):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            raise ValueError("Save CSV as UTF-8.") from None
        reader = csv.DictReader(io.StringIO(text))
        if not reader.fieldnames:
            raise ValueError("CSV headers are required.")
        headers = reader.fieldnames
        records = []
        for row in reader:
            if None in row:
                raise ValueError("CSV row has extra columns. Check delimiters and quotes.")
            records.append(row)
            if len(records) > 10000:
                raise ValueError("Maximum 10,000 rows per import.")
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(f.file_size for f in z.infolist()) > 64*1024*1024:
                    raise ValueError("XLSX exceeds the uncompressed size limit. Split the file.")
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                rows = book.active.iter_rows()
                first = next(rows, None)
                if not first:
                    raise ValueError("XLSX headers are required.")
                if any(c.data_type == "f" for c in first):
                    raise ValueError("Headers cannot contain Excel formulas.")
                headers = [c.value for c in first]
                records = []
                for cells in rows:
                    if any(c.data_type == "f" for c in cells):
                        raise ValueError("Excel formulas found. Export values before importing.")
                    records.append(dict(zip(headers, (c.value for c in cells))))
                    if len(records) > 10000:
                        raise ValueError("Maximum 10,000 rows per import.")
            finally:
                book.close()
        except (zipfile.BadZipFile, KeyError, OSError):
            raise ValueError("Cannot read the XLSX file.") from None
    else:
        raise ValueError("Choose a CSV or XLSX file.")
    names = [normalize(h) for h in headers if h is not None and str(h).strip()]
    if len(set(names)) != len(names):
        raise ValueError("Duplicate headers found.")
    expected = {normalize(a) for aliases in ALIASES.values() for a in aliases}
    if not set(names) & expected:
        raise ValueError("No SMILES or descriptor columns found. Use the template.")
    records = [row for row in records if any(v is not None and str(v).strip() for v in row.values())]
    if not records:
        raise ValueError("No data rows found.")
    return records


def uploaded_records(payload):
    if not isinstance(payload.get("filename"), str) or not isinstance(payload.get("content_base64"), str):
        raise ValueError("File name and content are required.")
    try:
        data = base64.b64decode(payload["content_base64"], validate=True)
    except (binascii.Error, ValueError):
        raise ValueError("Invalid file encoding.") from None
    return read_file(payload["filename"], data)


def batch_analyze(records):
    results = [analyze(row, i, draw=False) for i, row in enumerate(records, 1)]
    return {"results": results, "summary": {
        "total": len(results), "scored": sum(r["score_status"] == "ok" for r in results),
        "with_structure": sum(r["structure_status"] == "ok" for r in results),
        "recovered_structures": sum(r["structure_origin"] == "local_association" and r["structure_status"] == "ok" for r in results),
        "needs_input": sum(r["score_status"] != "ok" for r in results),
    }}


def results_csv(results):
    fields = ["row_id", "compound_id", "input_smiles", "canonical_smiles", "formula",
              "bbb_score_input", "bbb_score", "score_status", "structure_status",
              "descriptor_source", "pka_source", *FIELDS, "score_error", "structure_error", "warnings"]
    fields += ["structure_origin", "structure_source", "structure_source_url", "structure_source_locator", "structure_match_method"]
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=fields)
    writer.writeheader()
    for result in results:
        row = {k: result.get(k) for k in fields}
        row.update({k: result.get("descriptors", {}).get(k) for k in FIELDS})
        row.update({k: (result.get("structure") or {}).get(k) for k in ("canonical_smiles", "formula")})
        row["warnings"] = "; ".join(result["warnings"])
        # Prevent user-controlled text becoming executable spreadsheet formulas.
        for k, v in row.items():
            if isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")):
                row[k] = "'" + v
        writer.writerow(row)
    return stream.getvalue()
